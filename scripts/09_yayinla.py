"""Temiz katmanları tek veritabanına, Excel'e ve rapora yazar.

Çıktı : data/processed/istanbul72.duckdb   (tüm tablolar + temizleme_logu, kalite_kontrol, veri_sozlugu, kaynaklar)
        data/processed/*.parquet           (GeoParquet katmanları; web ve modelleme bunları okur)
        data/processed/istanbul72_veri.xlsx (gezmek için; biçimli, filtreli)
        docs/VERI_RAPORU.md                 (her çalıştırmada yeniden üretilir; elle düzenlenmez)
"""
from __future__ import annotations

import json
import shutil
from datetime import date

import duckdb
import geopandas as gpd
import pandas as pd
from kaynaklar import KAYNAKLAR
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from ortak import DOCS, INTERIM, LOG_DOSYA, PROCESSED, QC_DOSYA

print("09_yayinla")
KATMAN = {"hastane": "hastane", "itfaiye": "itfaiye", "toplanma_alani": "toplanma_alani", "acil_yol": "acil_yol",
          "sanat_yapisi": "sanat_yapisi", "yol_kenar": "yol_kenar", "yol_dugum": "yol_dugum",
          "silivri2025_istasyon": "silivri2025_istasyon", "silivri2025_dyfi": "silivri2025_dyfi", "senaryo_disi_bolge": "senaryo_disi_bolge"}
for ad in KATMAN:
    shutil.copy(INTERIM / f"{ad}.parquet", PROCESSED / f"{ad}.parquet")
shutil.copy(INTERIM / "yol_agi.graphml", PROCESSED / "yol_agi.graphml")

log = pd.DataFrame([json.loads(s) for s in LOG_DOSYA.read_text(encoding="utf-8").splitlines() if s])
qc = pd.DataFrame([json.loads(s) for s in QC_DOSYA.read_text(encoding="utf-8").splitlines() if s])
sira = sorted(log.adim.unique())
log = log.sort_values(["adim", "zaman"], key=lambda s: s.map({a: i for i, a in enumerate(sira)}) if s.name == "adim" else s).reset_index(drop=True)
log.insert(0, "no", range(1, len(log) + 1))
qc = qc.sort_values("adim", kind="stable").reset_index(drop=True)
soz = pd.read_parquet(INTERIM / "veri_sozlugu.parquet")
kay = pd.DataFrame(KAYNAKLAR)
mah = gpd.read_parquet(PROCESSED / "mahalle.parquet")

# ---------------- DuckDB ----------------
db = PROCESSED / "istanbul72.duckdb"
db.unlink(missing_ok=True)
con = duckdb.connect(str(db))
con.sql("INSTALL spatial; LOAD spatial;")
for ad in ["mahalle", *KATMAN]:
    yol = PROCESSED / f"{ad}.parquet"
    con.sql(f"CREATE TABLE {ad} AS SELECT * FROM read_parquet('{yol}')")
for ad, df in (("temizleme_logu", log), ("kalite_kontrol", qc), ("veri_sozlugu", soz), ("kaynaklar", kay)):
    con.register("_df", df); con.sql(f"CREATE TABLE {ad} AS SELECT * FROM _df"); con.unregister("_df")
con.sql("""
CREATE VIEW v_ilce_ozet AS
SELECT ilce, count(*) AS mahalle, sum(nufus_2025) AS nufus, sum(bina_toplam) AS bina_2017,
       sum(bina_cok_agir + bina_agir) AS agir_ustu_hasar, round(sum(bina_cok_agir + bina_agir) / sum(bina_toplam), 4) AS agir_ustu_oran,
       sum(can_kaybi) AS can_kaybi, round(1000.0 * sum(can_kaybi) / sum(nufus_2025), 3) AS can_kaybi_binde,
       sum(barinma_hane) AS barinma_hane, round(avg(vs30), 0) AS vs30_ort, round(min(fay_mesafe_km), 1) AS fay_min_km,
       sum(hastane_yatak_tahmini) AS yatak_tahmini, any_value(sosyal_gorebilirlik_ilce) AS sosyal_gorebilirlik
FROM mahalle GROUP BY ilce ORDER BY can_kaybi DESC""")
tablolar = con.sql("SELECT table_name, estimated_size FROM duckdb_tables() ORDER BY table_name").df()
ilce = con.sql("SELECT * FROM v_ilce_ozet").df()
con.close()
print(f"  duckdb: {len(tablolar)} tablo + v_ilce_ozet")

# ---------------- Excel ----------------
xl = PROCESSED / "istanbul72_veri.xlsx"
TURRENK = {"GÖZLEM": "D9EAD3", "MODEL": "FCE5CD", "TÜRETİLMİŞ": "DDE7F0", "VARSAYIM": "F4CCCC"}
oku = pd.DataFrame({"Sayfa": ["mahalle", "ilce_ozet", "hastane", "itfaiye", "temizleme_logu", "kalite_kontrol", "veri_sozlugu", "kaynaklar"],
                    "İçerik": ["959 mahalle × tüm değişkenler (geometri hariç)", "İlçe toplamları (v_ilce_ozet)", "377 hastane, tahmini yatak ve acil servis",
                               "136 itfaiye birimi", "Her temizleme kararı: sorun, etkilenen kayıt, karar, gerekçe", "Otomatik kontroller ve sonuçları",
                               "Sütun açıklaması, birim, kaynak, veri türü (renkli)", "Veri kaynakları, lisans, erişim tarihi"]})
with pd.ExcelWriter(xl, engine="openpyxl") as w:
    oku.to_excel(w, sheet_name="oku_beni", index=False)
    mah.drop(columns="geometry").to_excel(w, sheet_name="mahalle", index=False)
    ilce.to_excel(w, sheet_name="ilce_ozet", index=False)
    pd.read_parquet(PROCESSED / "hastane.parquet").drop(columns="geometry").to_excel(w, sheet_name="hastane", index=False)
    pd.read_parquet(PROCESSED / "itfaiye.parquet").drop(columns="geometry").to_excel(w, sheet_name="itfaiye", index=False)
    log.drop(columns="zaman").to_excel(w, sheet_name="temizleme_logu", index=False)
    qc.to_excel(w, sheet_name="kalite_kontrol", index=False)
    soz.to_excel(w, sheet_name="veri_sozlugu", index=False)
    kay.to_excel(w, sheet_name="kaynaklar", index=False)
    for ws in w.book.worksheets:
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2B2B2B")
        for i, col in enumerate(ws.columns, 1):
            uz = max(len(str(c.value or "")) for c in list(col)[:200])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, uz + 2), 70)
        if ws.title in ("temizleme_logu", "veri_sozlugu", "kaynaklar"):
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws = w.book["veri_sozlugu"]
    for row in ws.iter_rows(min_row=2):
        renk = TURRENK.get(row[4].value)
        if renk:
            for c in row:
                c.fill = PatternFill("solid", fgColor=renk)
    ws = w.book["kalite_kontrol"]
    for row in ws.iter_rows(min_row=2):
        row[2].fill = PatternFill("solid", fgColor="D9EAD3" if row[2].value == "GEÇTİ" else "F4CCCC")
print(f"  excel: {xl.name}")

# ---------------- rapor ----------------
def md_tablo(df: pd.DataFrame) -> str:
    satir = ["| " + " | ".join(map(str, df.columns)) + " |", "|" + "---|" * len(df.columns)]
    satir += ["| " + " | ".join(str(v).replace("|", "/").replace("\n", " ") for v in r) + " |" for r in df.itertuples(index=False)]
    return "\n".join(satir)

n_gecti = int((qc.sonuc == "GEÇTİ").sum())
tur_say = soz.veri_turu.value_counts()
r = [f"# Veri hazırlama raporu\n\n_Otomatik üretildi: `scripts/09_yayinla.py`, {date.today().isoformat()}. Elle düzenlemeyin; betikleri değiştirip hattı yeniden çalıştırın._\n",
     "## Özet\n",
     "- **959 mahalle**, 39 ilçe; her mahallede senaryo, bina stoğu, 2025 nüfusu, güncel sınır, zemin, fay mesafesi, yol ağı erişim düğümü ve tesis sayıları var.",
     f"- **{len(log)} temizleme kararı** kayıtlı (sorun → etkilenen kayıt → karar → gerekçe).",
     f"- **{len(qc)} otomatik kontrol**, {n_gecti} geçti, {len(qc) - n_gecti} kaldı.",
     f"- Mahalle tablosunun {len(soz)} sütunu veri türüne göre: gözlem {tur_say.get('GÖZLEM', 0)}, başka bir modelin çıktısı "
     f"{tur_say.get('MODEL', 0)}, türetilmiş {tur_say.get('TÜRETİLMİŞ', 0)}, varsayım {tur_say.get('VARSAYIM', 0)}.",
     f"- Yol ağı: {len(pd.read_parquet(PROCESSED / 'yol_dugum.parquet', columns=['osmid'])):,} düğüm, {len(pd.read_parquet(PROCESSED / 'yol_kenar.parquet', columns=['u'])):,} yönlü kenar.\n",
     "## Hat\n",
     "```mermaid\nflowchart LR\n  A[Geofabrik Türkiye OSM 650 MB] --> K[01 İstanbul kesimi]\n  B[Overture 1,3 M bina] --> K\n  C[USGS Vs30 1 GB] --> K\n"
     "  D[İBB senaryo + bina 2017] --> T[02 mahalle tablosu]\n  E[TÜİK 2025 nüfus] --> T\n  F[İBB sosyal göstergeler] --> T\n"
     "  K --> S[03 sınır eşleştirme]\n  T --> S\n  G[İBB muhtarlık noktaları] --> S\n  K --> Y[05 yol ağı]\n  H[İBB acil yol + sanat yapıları] --> Y\n"
     "  K --> TS[04 tesisler]\n  I[İBB sağlık + itfaiye + yeşil alan] --> TS\n  S --> Z[06 zemin + bina]\n  Y --> Z\n  J[USGS Silivri 2025 kayıtları] --> O[07 gözlem]\n"
     "  S --> B8[08 birleştir]\n  Z --> B8\n  O --> B8\n  TS --> B8\n  Y --> B8\n  B8 --> P[(09 DuckDB + Excel + GeoParquet)]\n```\n",
     "## Kaynaklar\n", md_tablo(kay[["katman", "kaynak", "lisans", "erisim"]]), "\n",
     "## Temizleme kararları\n"]
for adim, g in log.groupby("adim", sort=False):
    r.append(f"### {adim}\n")
    r.append(md_tablo(g[["no", "katman", "sorun", "etkilenen", "karar", "gerekce"]].rename(columns={"gerekce": "gerekçe"})))
    r.append("")
r += ["## Otomatik kontroller\n", md_tablo(qc[["adim", "kontrol", "sonuc", "detay"]]), "",
      "## Bilinen sınırlılıklar\n",
      "- İBB senaryo sütunları **model çıktısıdır**, gözlem değildir. Bu proje gerçek bir depremi tahmin etmez; senaryonun neyi neye bağladığını açıklar ve üstüne bir müdahale simülasyonu kurar.",
      "- Bina stoğu 2017, nüfus 2025. Overture bina sayısıyla mahalle sırası ρ = 0,83 tutarlı; tarihî yarımadada Overture blokları birleşik çizdiği için az sayar.",
      "- Vs30 (~900 m, topoğrafya vekili) kıyı dolgularını ve alüvyonu yakalayamaz; İstanbul'da E sınıfı hiç çıkmıyor, gerçekte var.",
      "- Hastane yatakları tesis bazında açık veride yok; il toplamı kat alanına göre dağıtıldı (OSM'deki 8 yatak etiketiyle ρ = 0,60, zayıf doğrulama).",
      "- Resmî AFAD toplanma alanı listesi indirilebilir değil; açık alan katmanı 'potansiyel' olarak etiketli. Uygun alan sayısı (2.960) AFAD'ın açıkladığı 3.021 ile aynı mertebede.",
      "- Sosyal görebilirlik yalnızca ilçe düzeyinde (39 değer); mahalle düzeyi yoksunluk vekili olarak sosyal yardım oranı kullanılabilir.\n",
      "## Dosya haritası\n",
      "| Yol | Ne |\n|---|---|\n| `data/raw/` | Kaynaktan indirildiği gibi, dokunulmaz (repoya girmez) |\n"
      "| `data/interim/` | Ara ürünler: İstanbul kesimleri, eşleştirme, log/QC JSONL |\n| `data/processed/istanbul72.duckdb` | Tüm temiz tablolar + log + QC + sözlük + kaynaklar |\n"
      "| `data/processed/*.parquet` | GeoParquet katmanları (mahalle, hastane, yol_kenar …) |\n| `data/processed/istanbul72_veri.xlsx` | Gezmek için Excel |\n"
      "| `data/interim/eslesme_inceleme.csv` | 'ad+muhtarlık' dışındaki her sınır eşleşmesi, elle göz atmak için |\n| `scripts/01…09` | Hat; `python scripts/hat.py` hepsini sırayla çalıştırır |\n"]
(DOCS / "VERI_RAPORU.md").write_text("\n".join(r), encoding="utf-8")
print(f"  rapor: docs/VERI_RAPORU.md ({len(log)} karar, {len(qc)} kontrol)")
